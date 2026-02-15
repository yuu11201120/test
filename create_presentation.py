#!/usr/bin/env python3
"""KDDIの将来性に関するパワーポイント資料を作成するスクリプト"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# カラーパレット（KDDIブランドカラーベース）
KDDI_ORANGE = RGBColor(0xFF, 0x66, 0x00)
KDDI_DARK = RGBColor(0x33, 0x33, 0x33)
KDDI_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
KDDI_LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)
KDDI_MEDIUM_GRAY = RGBColor(0x99, 0x99, 0x99)
ACCENT_BLUE = RGBColor(0x00, 0x78, 0xD4)
ACCENT_GREEN = RGBColor(0x00, 0xA5, 0x5A)
ACCENT_PURPLE = RGBColor(0x6B, 0x3F, 0xA0)
ACCENT_RED = RGBColor(0xE0, 0x3C, 0x31)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def add_background(slide, color):
    """スライドの背景色を設定"""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, left, top, width, height, fill_color, border_color=None):
    """色付きの四角形を追加"""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_rounded_rect(slide, left, top, width, height, fill_color):
    """角丸四角形を追加"""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape


def add_textbox(slide, left, top, width, height, text, font_size=18,
                font_color=KDDI_DARK, bold=False, alignment=PP_ALIGN.LEFT,
                font_name="Meiryo"):
    """テキストボックスを追加"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = font_color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_multi_text(slide, left, top, width, height, lines, font_size=16,
                   font_color=KDDI_DARK, bold=False, line_spacing=1.5,
                   font_name="Meiryo", alignment=PP_ALIGN.LEFT):
    """複数行テキストを追加"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(font_size)
        p.font.color.rgb = font_color
        p.font.bold = bold
        p.font.name = font_name
        p.alignment = alignment
        p.space_after = Pt(font_size * (line_spacing - 1))
    return txBox


def add_bullet_points(slide, left, top, width, height, items, font_size=15,
                      font_color=KDDI_DARK, bullet_char="\u2022",
                      font_name="Meiryo", line_spacing=1.4):
    """箇条書きテキストを追加"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"{bullet_char} {item}"
        p.font.size = Pt(font_size)
        p.font.color.rgb = font_color
        p.font.name = font_name
        p.space_after = Pt(font_size * (line_spacing - 1) * 0.8)
    return txBox


def add_section_header(slide, text):
    """セクションヘッダーバーを追加"""
    # ヘッダーバー
    add_shape(slide, Inches(0), Inches(0), SLIDE_W, Inches(1.1), KDDI_ORANGE)
    # タイトルテキスト
    add_textbox(slide, Inches(0.8), Inches(0.15), Inches(11), Inches(0.8),
                text, font_size=30, font_color=KDDI_WHITE, bold=True)
    # 下線アクセント
    add_shape(slide, Inches(0.8), Inches(1.0), Inches(2), Inches(0.05), KDDI_WHITE)


def add_page_number(slide, num, total):
    """ページ番号を追加"""
    add_textbox(slide, Inches(12.0), Inches(7.0), Inches(1.2), Inches(0.4),
                f"{num} / {total}", font_size=10, font_color=KDDI_MEDIUM_GRAY,
                alignment=PP_ALIGN.RIGHT)


# =====================================================
# スライド1: 表紙
# =====================================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # blank
add_background(slide1, KDDI_WHITE)

# 上部のオレンジバー
add_shape(slide1, Inches(0), Inches(0), SLIDE_W, Inches(0.15), KDDI_ORANGE)

# メインビジュアルエリア
add_shape(slide1, Inches(0), Inches(1.5), SLIDE_W, Inches(3.5), KDDI_DARK)

# タイトル
add_textbox(slide1, Inches(1.5), Inches(1.8), Inches(10), Inches(1.2),
            "KDDIの将来性", font_size=48, font_color=KDDI_WHITE, bold=True,
            alignment=PP_ALIGN.CENTER)

# サブタイトル
add_textbox(slide1, Inches(1.5), Inches(3.0), Inches(10), Inches(0.8),
            "〜 通信を超えた成長戦略と未来への挑戦 〜",
            font_size=24, font_color=KDDI_ORANGE, bold=False,
            alignment=PP_ALIGN.CENTER)

# 区切り線
add_shape(slide1, Inches(5.5), Inches(3.8), Inches(2.3), Inches(0.03), KDDI_ORANGE)

# 日付
add_textbox(slide1, Inches(1.5), Inches(4.1), Inches(10), Inches(0.5),
            "2026年2月", font_size=18, font_color=KDDI_MEDIUM_GRAY,
            alignment=PP_ALIGN.CENTER)

# 下部の説明テキスト
add_multi_text(slide1, Inches(2.5), Inches(5.5), Inches(8), Inches(1.5),
               ["5G・AI・DX・衛星通信・ローソン提携 ── 多角的成長戦略の全貌"],
               font_size=16, font_color=KDDI_DARK, alignment=PP_ALIGN.CENTER)

# 下部のオレンジバー
add_shape(slide1, Inches(0), Inches(7.35), SLIDE_W, Inches(0.15), KDDI_ORANGE)


# =====================================================
# スライド2: 目次
# =====================================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide2, KDDI_WHITE)
add_section_header(slide2, "目次")

toc_items = [
    ("01", "KDDIの企業概要と現在の業績", "事業規模・収益構造・直近決算"),
    ("02", "中期経営戦略「新サテライトグロース戦略」", "コア事業とOrbit1/Orbit2の成長領域"),
    ("03", "5G・Beyond 5G/6G戦略", "次世代通信インフラへの取り組み"),
    ("04", "AI・DX事業の成長", "WAKONX・データセンター・法人DX"),
    ("05", "宇宙事業 ── au Starlink Direct", "衛星直接通信と月面5G構想"),
    ("06", "ローソン提携によるシナジー", "リアル×デジタル×グリーンの融合"),
    ("07", "金融・エネルギー事業", "非通信収益の多角化戦略"),
    ("08", "リスク要因と課題", "競争環境・規制・子会社問題"),
    ("09", "まとめ ── KDDIの将来展望", "成長シナリオと投資判断のポイント"),
]

for i, (num, title, desc) in enumerate(toc_items):
    y = Inches(1.5) + Inches(i * 0.62)
    # 番号
    add_textbox(slide2, Inches(1.2), y, Inches(0.8), Inches(0.5),
                num, font_size=22, font_color=KDDI_ORANGE, bold=True,
                alignment=PP_ALIGN.CENTER)
    # 区切り線
    add_shape(slide2, Inches(2.1), y + Inches(0.05), Inches(0.03), Inches(0.4), KDDI_ORANGE)
    # タイトル
    add_textbox(slide2, Inches(2.4), y - Inches(0.05), Inches(6), Inches(0.4),
                title, font_size=18, font_color=KDDI_DARK, bold=True)
    # 説明
    add_textbox(slide2, Inches(2.4), y + Inches(0.25), Inches(6), Inches(0.3),
                desc, font_size=12, font_color=KDDI_MEDIUM_GRAY)

add_page_number(slide2, 2, 10)


# =====================================================
# スライド3: 企業概要と業績
# =====================================================
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide3, KDDI_WHITE)
add_section_header(slide3, "01  KDDIの企業概要と現在の業績")

# 左カラム: 企業概要
add_shape(slide3, Inches(0.6), Inches(1.4), Inches(5.8), Inches(0.5), KDDI_LIGHT_GRAY)
add_textbox(slide3, Inches(0.8), Inches(1.45), Inches(5), Inches(0.4),
            "企業概要", font_size=18, font_color=KDDI_DARK, bold=True)

overview_items = [
    "設立: 2000年（DDI・KDD・IDO の合併）",
    "本社: 東京都千代田区飯田橋 → 2025年 高輪ゲートウェイへ移転",
    "従業員数: 約62,000名（連結）",
    "事業: 通信（au/UQ/povo）、DX、金融、エネルギー等",
    "時価総額: 約10兆円規模（通信業界2位）",
]
add_bullet_points(slide3, Inches(0.8), Inches(2.0), Inches(5.5), Inches(3.5),
                  overview_items, font_size=14)

# 右カラム: 業績ハイライト
add_shape(slide3, Inches(6.8), Inches(1.4), Inches(5.8), Inches(0.5), KDDI_LIGHT_GRAY)
add_textbox(slide3, Inches(7.0), Inches(1.45), Inches(5), Inches(0.4),
            "2025年3月期 業績ハイライト", font_size=18, font_color=KDDI_DARK, bold=True)

# 業績データカード
metrics = [
    ("売上高", "5兆9,180億円", "+2.8%"),
    ("営業利益", "1兆1,187億円", "+16.3%"),
    ("決済・金融取扱高", "21.4兆円", "+3.4兆円"),
]

for i, (label, value, change) in enumerate(metrics):
    y = Inches(2.1) + Inches(i * 1.2)
    card = add_rounded_rect(slide3, Inches(7.0), y, Inches(5.3), Inches(1.0), KDDI_LIGHT_GRAY)
    add_textbox(slide3, Inches(7.3), y + Inches(0.05), Inches(3), Inches(0.4),
                label, font_size=13, font_color=KDDI_MEDIUM_GRAY, bold=True)
    add_textbox(slide3, Inches(7.3), y + Inches(0.35), Inches(3), Inches(0.5),
                value, font_size=22, font_color=KDDI_DARK, bold=True)
    add_textbox(slide3, Inches(10.5), y + Inches(0.35), Inches(1.5), Inches(0.5),
                change, font_size=18, font_color=ACCENT_GREEN, bold=True,
                alignment=PP_ALIGN.RIGHT)

# 2026年3月期 上期
add_shape(slide3, Inches(0.6), Inches(5.4), Inches(12.1), Inches(0.05), KDDI_ORANGE)
add_textbox(slide3, Inches(0.8), Inches(5.6), Inches(12), Inches(0.4),
            "2026年3月期 上期（2025年4月〜9月）: 売上高 2兆9,632億円（+3.8%）/ 営業利益 5,772億円（+0.7%）/ 当期利益 3,777億円（+7.6%）",
            font_size=14, font_color=KDDI_DARK)

add_textbox(slide3, Inches(0.8), Inches(6.1), Inches(12), Inches(0.4),
            "モバイルARPU: 4,460円（Q1比+120円）/ 解約率: 1.21%（改善傾向）",
            font_size=14, font_color=ACCENT_BLUE, bold=True)

add_page_number(slide3, 3, 10)


# =====================================================
# スライド4: 新サテライトグロース戦略
# =====================================================
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide4, KDDI_WHITE)
add_section_header(slide4, "02  中期経営戦略「新サテライトグロース戦略」")

# ビジョン
add_rounded_rect(slide4, Inches(0.8), Inches(1.4), Inches(11.7), Inches(0.9), KDDI_LIGHT_GRAY)
add_textbox(slide4, Inches(1.0), Inches(1.45), Inches(11.3), Inches(0.8),
            "KDDI VISION 2030:「つなぐチカラ」を進化させ、誰もが思いを実現できる社会をつくる。",
            font_size=17, font_color=KDDI_DARK, bold=True, alignment=PP_ALIGN.CENTER)

# コア事業
add_shape(slide4, Inches(4.5), Inches(2.6), Inches(4.3), Inches(1.2), KDDI_ORANGE)
add_textbox(slide4, Inches(4.7), Inches(2.65), Inches(3.9), Inches(0.4),
            "コア事業", font_size=16, font_color=KDDI_WHITE, bold=True,
            alignment=PP_ALIGN.CENTER)
add_textbox(slide4, Inches(4.7), Inches(3.0), Inches(3.9), Inches(0.7),
            "5G通信 × データドリブン × 生成AI",
            font_size=14, font_color=KDDI_WHITE, alignment=PP_ALIGN.CENTER)

# Orbit 1
add_rounded_rect(slide4, Inches(0.8), Inches(4.2), Inches(5.5), Inches(2.6), RGBColor(0xE8, 0xF0, 0xFE))
add_shape(slide4, Inches(0.8), Inches(4.2), Inches(5.5), Inches(0.5), ACCENT_BLUE)
add_textbox(slide4, Inches(1.0), Inches(4.22), Inches(5), Inches(0.45),
            "Orbit 1 ── 成長牽引領域", font_size=16, font_color=KDDI_WHITE, bold=True)

orbit1_items = [
    "DX（デジタルトランスフォーメーション）",
    "  → WAKONX基盤、IoT、データセンター拡充",
    "金融事業",
    "  → 決済・金融取扱高 21.4兆円、au PAY/じぶん銀行",
    "エネルギー事業",
    "  → auでんき、再生可能エネルギー推進",
]
add_bullet_points(slide4, Inches(1.0), Inches(4.8), Inches(5.0), Inches(2.0),
                  orbit1_items, font_size=12, bullet_char="▸")

# Orbit 2
add_rounded_rect(slide4, Inches(6.8), Inches(4.2), Inches(5.7), Inches(2.6), RGBColor(0xF3, 0xE8, 0xFD))
add_shape(slide4, Inches(6.8), Inches(4.2), Inches(5.7), Inches(0.5), ACCENT_PURPLE)
add_textbox(slide4, Inches(7.0), Inches(4.22), Inches(5), Inches(0.45),
            "Orbit 2 ── 将来成長領域", font_size=16, font_color=KDDI_WHITE, bold=True)

orbit2_items = [
    "モビリティ（自動運転・MaaS）",
    "宇宙事業（Starlink・月面5G）",
    "ヘルスケア（遠隔医療・健康管理）",
    "Web3・メタバース（αU等）",
    "スポーツ・エンターテインメント",
]
add_bullet_points(slide4, Inches(7.0), Inches(4.8), Inches(5.2), Inches(2.0),
                  orbit2_items, font_size=12, bullet_char="▸")

# 財務目標
add_textbox(slide4, Inches(0.8), Inches(7.0), Inches(12), Inches(0.4),
            "目標: 営業CF 2年間で3兆円規模 / 設備投資 1.3兆円 / 戦略投資 2,000億円 / 持続的増配・配当性向40%超",
            font_size=13, font_color=KDDI_DARK, bold=True)

add_page_number(slide4, 4, 10)


# =====================================================
# スライド5: 5G・Beyond 5G/6G
# =====================================================
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide5, KDDI_WHITE)
add_section_header(slide5, "03  5G・Beyond 5G/6G戦略")

# 左カラム: 5G現状
add_shape(slide5, Inches(0.6), Inches(1.4), Inches(5.8), Inches(0.5), ACCENT_BLUE)
add_textbox(slide5, Inches(0.8), Inches(1.43), Inches(5), Inches(0.4),
            "5G（現在〜2030年）", font_size=18, font_color=KDDI_WHITE, bold=True)

items_5g = [
    "Sub-6/ミリ波による高速大容量通信の全国展開",
    "5G SA（スタンドアロン）の本格化",
    "ネットワークスライシングによる用途別最適化",
    "インフラシェアリング推進でコスト効率化",
    "中期1,000億円規模のコスト効率化を目標",
    "モバイルARPU向上施策の推進",
]
add_bullet_points(slide5, Inches(0.8), Inches(2.1), Inches(5.3), Inches(3.5),
                  items_5g, font_size=14)

# 右カラム: Beyond 5G/6G
add_shape(slide5, Inches(6.8), Inches(1.4), Inches(5.8), Inches(0.5), ACCENT_PURPLE)
add_textbox(slide5, Inches(7.0), Inches(1.43), Inches(5), Inches(0.4),
            "Beyond 5G / 6G（2030年〜）", font_size=18, font_color=KDDI_WHITE, bold=True)

items_6g = [
    "「KDDI Accelerate 5.0」ビジョン",
    "7つの重点テクノロジー:",
    "  ネットワーク / セキュリティ / IoT",
    "  プラットフォーム / AI / XR / ロボティクス",
    "超高速暗号アルゴリズム「Rocca」（100Gbps超）",
    "点群圧縮技術（PCC）: データ量1/40に",
    "光無線融合技術（IFoF）: アンテナ小型化・省電力化",
    "ユーザーセントリックアーキテクチャへの転換",
]
add_bullet_points(slide5, Inches(7.0), Inches(2.1), Inches(5.3), Inches(4.0),
                  items_6g, font_size=13)

# ロードマップ
add_shape(slide5, Inches(0.6), Inches(5.8), Inches(12.1), Inches(0.05), KDDI_ORANGE)

# タイムラインバー
phases = [
    (Inches(0.8), Inches(3.5), "2024-2026", "5G SA本格展開\nARPU回復期", ACCENT_BLUE),
    (Inches(4.5), Inches(3.5), "2026-2030", "5G Advanced\n非通信収益拡大", RGBColor(0x00, 0x96, 0xC7)),
    (Inches(8.2), Inches(3.5), "2030〜", "Beyond 5G/6G\n社会プラットフォーマー", ACCENT_PURPLE),
]
for x, w, label, desc, color in phases:
    add_rounded_rect(slide5, x, Inches(6.0), w, Inches(1.1), color)
    add_textbox(slide5, x + Inches(0.15), Inches(6.05), w - Inches(0.3), Inches(0.4),
                label, font_size=13, font_color=KDDI_WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    add_textbox(slide5, x + Inches(0.15), Inches(6.4), w - Inches(0.3), Inches(0.6),
                desc, font_size=11, font_color=KDDI_WHITE, alignment=PP_ALIGN.CENTER)

add_page_number(slide5, 5, 10)


# =====================================================
# スライド6: AI・DX事業
# =====================================================
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide6, KDDI_WHITE)
add_section_header(slide6, "04  AI・DX事業の成長")

# WAKONX
add_rounded_rect(slide6, Inches(0.6), Inches(1.5), Inches(6.0), Inches(2.5), RGBColor(0xE8, 0xF0, 0xFE))
add_textbox(slide6, Inches(0.8), Inches(1.55), Inches(5.5), Inches(0.5),
            "WAKONX（ワコンクロス）", font_size=20, font_color=ACCENT_BLUE, bold=True)
wakonx_items = [
    "AI時代のビジネスプラットフォーム",
    "AI・データ基盤・ネットワーク運用をワンストップ提供",
    "社会課題・業界課題の解決を加速",
    "法人事業のID拡大と中小企業層アプローチ強化",
]
add_bullet_points(slide6, Inches(0.8), Inches(2.2), Inches(5.5), Inches(1.8),
                  wakonx_items, font_size=13, bullet_char="▸")

# データセンター・AI投資
add_rounded_rect(slide6, Inches(6.9), Inches(1.5), Inches(5.8), Inches(2.5), RGBColor(0xFE, 0xF0, 0xE8))
add_textbox(slide6, Inches(7.1), Inches(1.55), Inches(5.3), Inches(0.5),
            "AIインフラ・データセンター投資", font_size=20, font_color=KDDI_ORANGE, bold=True)
dc_items = [
    "NVIDIA GB200 NVL72システム導入（2026年1月〜）",
    "ロンドンに6番目の施設（約600億円投資）",
    "DC電力容量 1.6倍に拡大",
    "生成AIの社会実装を推進",
]
add_bullet_points(slide6, Inches(7.1), Inches(2.2), Inches(5.3), Inches(1.8),
                  dc_items, font_size=13, bullet_char="▸")

# ビジネスセグメント業績
add_shape(slide6, Inches(0.6), Inches(4.3), Inches(12.1), Inches(0.5), KDDI_DARK)
add_textbox(slide6, Inches(0.8), Inches(4.33), Inches(11.5), Inches(0.4),
            "ビジネスセグメント業績目標", font_size=18, font_color=KDDI_WHITE, bold=True)

biz_metrics = [
    ("売上高", "1兆5,900億円", "前年比+13.6%"),
    ("営業利益", "2,720億円", "前年比+16.7%"),
    ("重点領域", "IoT・DC等グロース領域", "二桁成長を牽引"),
]

for i, (label, value, note) in enumerate(biz_metrics):
    x = Inches(0.8) + Inches(i * 4.1)
    add_rounded_rect(slide6, x, Inches(5.0), Inches(3.7), Inches(1.2), KDDI_LIGHT_GRAY)
    add_textbox(slide6, x + Inches(0.2), Inches(5.05), Inches(3.3), Inches(0.3),
                label, font_size=12, font_color=KDDI_MEDIUM_GRAY, bold=True)
    add_textbox(slide6, x + Inches(0.2), Inches(5.35), Inches(3.3), Inches(0.4),
                value, font_size=20, font_color=KDDI_DARK, bold=True)
    add_textbox(slide6, x + Inches(0.2), Inches(5.75), Inches(3.3), Inches(0.3),
                note, font_size=12, font_color=ACCENT_GREEN, bold=True)

# DX銘柄
add_textbox(slide6, Inches(0.8), Inches(6.5), Inches(12), Inches(0.4),
            "「DX銘柄2025」に選定 ── 「DX実現能力の高さでは、確実にトップランナーの1社」（DX銘柄選考委員）",
            font_size=14, font_color=ACCENT_BLUE, bold=True)

add_page_number(slide6, 6, 10)


# =====================================================
# スライド7: 宇宙事業
# =====================================================
slide7 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide7, KDDI_WHITE)
add_section_header(slide7, "05  宇宙事業 ── au Starlink Direct")

# メインコンテンツ
add_rounded_rect(slide7, Inches(0.6), Inches(1.5), Inches(7.5), Inches(5.3), RGBColor(0xE8, 0xF0, 0xFE))

add_textbox(slide7, Inches(0.8), Inches(1.6), Inches(7.0), Inches(0.5),
            "au Starlink Direct（2025年4月サービス開始）",
            font_size=20, font_color=ACCENT_BLUE, bold=True)

starlink_items = [
    "国内キャリア初・アジア初の衛星⇔スマホ直接通信サービス",
    "「空が見えれば、どこでもつながる」── 圏外エリアを解消",
    "対応機種: Android/iPhone 50機種・約600万台",
    "開始1カ月で約20万人が利用（GW中は1日4万人）",
    "申込不要・当面無料で提供",
    "SMS/RCS/iMessage/緊急地震速報に対応",
    "2025年夏からデータ通信にも対応予定",
    "高度340kmの低軌道衛星（約600機）と直接通信",
    "60年の衛星通信ノウハウがSpaceXとの提携を実現",
]
add_bullet_points(slide7, Inches(0.8), Inches(2.2), Inches(7.0), Inches(4.5),
                  starlink_items, font_size=13, bullet_char="▸", line_spacing=1.3)

# 月面5G構想
add_rounded_rect(slide7, Inches(8.4), Inches(1.5), Inches(4.3), Inches(2.5), RGBColor(0x2D, 0x2D, 0x4E))
add_textbox(slide7, Inches(8.6), Inches(1.6), Inches(3.9), Inches(0.5),
            "月面5G構想", font_size=18, font_color=KDDI_ORANGE, bold=True)
moon_items = [
    "月面機器間の5G通信",
    "月面⇔衛星⇔地球の通信網",
    "JAXA・GITAIとの共同研究",
    "2025年2月〜具体活動開始",
]
add_bullet_points(slide7, Inches(8.6), Inches(2.2), Inches(3.9), Inches(1.5),
                  moon_items, font_size=12, font_color=KDDI_WHITE, bullet_char="▸")

# 法人向け
add_rounded_rect(slide7, Inches(8.4), Inches(4.3), Inches(4.3), Inches(2.5), RGBColor(0xFE, 0xF0, 0xE8))
add_textbox(slide7, Inches(8.6), Inches(4.4), Inches(3.9), Inches(0.5),
            "法人向け展開", font_size=18, font_color=KDDI_ORANGE, bold=True)
biz_starlink = [
    "Starlink Businessによる",
    "閉域ネットワーク実証成功",
    "災害時のBCP対策需要",
    "僻地・海上での安定通信",
]
add_bullet_points(slide7, Inches(8.6), Inches(5.0), Inches(3.9), Inches(1.5),
                  biz_starlink, font_size=12, bullet_char="▸")

add_page_number(slide7, 7, 10)


# =====================================================
# スライド8: ローソン提携
# =====================================================
slide8 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide8, KDDI_WHITE)
add_section_header(slide8, "06  ローソン提携によるシナジー")

# 提携概要
add_rounded_rect(slide8, Inches(0.6), Inches(1.5), Inches(12.1), Inches(1.2), KDDI_LIGHT_GRAY)
add_textbox(slide8, Inches(0.8), Inches(1.55), Inches(11.5), Inches(0.4),
            "三菱商事・KDDI・ローソン 資本業務提携（2024年2月締結 → 8月共同経営開始）",
            font_size=17, font_color=KDDI_DARK, bold=True)
add_textbox(slide8, Inches(0.8), Inches(1.95), Inches(11.5), Inches(0.5),
            "テーマ: 「リアル × デジタル × グリーン」の融合 ── 三菱商事・KDDIがローソン株を50%ずつ保有",
            font_size=14, font_color=KDDI_DARK)

# シナジー領域
synergy_areas = [
    ("未来のコンビニ", ACCENT_BLUE,
     ["AI活用スマホレジ・デジタルサイネージ",
      "フライド商品自動調理器・飲料自動補充ロボ",
      "アバターによる遠隔接客",
      "高輪ゲートウェイに実験店舗（2025年6月）"]),
    ("データ基盤統合", ACCENT_GREEN,
     ["KDDIとローソンの会員情報連携",
      "国内最大級の顧客データ基盤構築",
      "パーソナライズされた顧客体験",
      "Ponta会員基盤の活用拡大"]),
    ("地域課題解決", ACCENT_PURPLE,
     ["防災・災害対処協定の締結",
      "Starlink通信環境の店舗活用",
      "ドローン周辺パトロール",
      "オンデマンド乗合交通連携"]),
]

for i, (title, color, items) in enumerate(synergy_areas):
    x = Inches(0.6) + Inches(i * 4.1)
    add_shape(slide8, x, Inches(3.0), Inches(3.8), Inches(0.5), color)
    add_textbox(slide8, x + Inches(0.15), Inches(3.03), Inches(3.5), Inches(0.4),
                title, font_size=15, font_color=KDDI_WHITE, bold=True,
                alignment=PP_ALIGN.CENTER)
    add_bullet_points(slide8, x + Inches(0.15), Inches(3.6), Inches(3.5), Inches(2.5),
                      items, font_size=12, bullet_char="▸", line_spacing=1.3)

# 課題
add_shape(slide8, Inches(0.6), Inches(5.8), Inches(12.1), Inches(0.05), ACCENT_RED)
add_textbox(slide8, Inches(0.8), Inches(6.0), Inches(11.5), Inches(0.4),
            "課題: 「キャリア色」を出しすぎると他社契約者の離反リスク / セブンとの日販差（約12万円）をどう縮めるか",
            font_size=14, font_color=ACCENT_RED, bold=True)
add_textbox(slide8, Inches(0.8), Inches(6.4), Inches(11.5), Inches(0.4),
            "展望: 株価は買収後20%下落→回復。「資本コスト分のリターンは確保できる」との評価に転換",
            font_size=14, font_color=ACCENT_GREEN)

add_page_number(slide8, 8, 10)


# =====================================================
# スライド9: 金融・エネルギー＋リスク
# =====================================================
slide9 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide9, KDDI_WHITE)
add_section_header(slide9, "07-08  金融・エネルギー事業 / リスク要因と課題")

# 左半分: 金融・エネルギー
add_shape(slide9, Inches(0.6), Inches(1.4), Inches(5.8), Inches(0.5), ACCENT_GREEN)
add_textbox(slide9, Inches(0.8), Inches(1.43), Inches(5), Inches(0.4),
            "金融・エネルギー事業", font_size=18, font_color=KDDI_WHITE, bold=True)

finance_items = [
    "auじぶん銀行・auカブコム証券・au PAY",
    "決済・金融取扱高: 21.4兆円（前年+3.4兆円）",
    "通信×金融の好循環モデル確立",
    "金融事業で20%成長を目指す方針",
    "金利環境変化への対応（預貸率重視へ転換）",
    "",
    "auでんき・再生可能エネルギー事業",
    "サステナビリティファイナンスの推進",
    "下期注力領域で DX・金融・エネルギー・ローソン合計 +約300億円目標",
]
add_bullet_points(slide9, Inches(0.8), Inches(2.1), Inches(5.3), Inches(4.5),
                  finance_items, font_size=13, bullet_char="▸", line_spacing=1.25)

# 右半分: リスク
add_shape(slide9, Inches(6.8), Inches(1.4), Inches(5.8), Inches(0.5), ACCENT_RED)
add_textbox(slide9, Inches(7.0), Inches(1.43), Inches(5), Inches(0.4),
            "リスク要因と課題", font_size=18, font_color=KDDI_WHITE, bold=True)

risk_items = [
    ("競争環境の激化", [
        "楽天モバイルの衛星直接通信（2026年開始予定）",
        "ドコモのHAPS（空飛ぶ基地局、2026年商用化）",
        "格安プラン競争（povo vs ahamo vs LINEMO）",
    ]),
    ("規制・政策リスク", [
        "通信料値下げ圧力の継続",
        "電気通信事業法改正の影響",
    ]),
    ("子会社不適切取引", [
        "ビッグローブ・ジー・プラン広告事業で架空取引",
        "特別調査委員会が調査中（2026年3月末報告予定）",
    ]),
    ("金融事業環境", [
        "金利上昇局面での預金調達競争激化",
    ]),
]

y_pos = Inches(2.1)
for category, items in risk_items:
    add_textbox(slide9, Inches(7.0), y_pos, Inches(5.3), Inches(0.35),
                category, font_size=14, font_color=ACCENT_RED, bold=True)
    y_pos += Inches(0.35)
    for item in items:
        add_textbox(slide9, Inches(7.3), y_pos, Inches(5.0), Inches(0.3),
                    f"- {item}", font_size=11, font_color=KDDI_DARK)
        y_pos += Inches(0.28)
    y_pos += Inches(0.1)

add_page_number(slide9, 9, 10)


# =====================================================
# スライド10: まとめ
# =====================================================
slide10 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide10, KDDI_WHITE)
add_section_header(slide10, "09  まとめ ── KDDIの将来展望")

# 成長ドライバー
add_textbox(slide10, Inches(0.8), Inches(1.5), Inches(12), Inches(0.5),
            "KDDIの5つの成長ドライバー", font_size=22, font_color=KDDI_DARK, bold=True)

drivers = [
    ("5G/6G", "次世代通信\nインフラ", ACCENT_BLUE),
    ("AI・DX", "WAKONX基盤\n法人事業拡大", KDDI_ORANGE),
    ("宇宙", "Starlink直接通信\n月面5G", ACCENT_PURPLE),
    ("ローソン", "リアル×デジタル\n13,000+店舗", ACCENT_GREEN),
    ("金融", "21兆円取扱高\n通信×金融融合", RGBColor(0xC0, 0x7A, 0x00)),
]

for i, (title, desc, color) in enumerate(drivers):
    x = Inches(0.6) + Inches(i * 2.5)
    add_rounded_rect(slide10, x, Inches(2.2), Inches(2.2), Inches(2.0), color)
    add_textbox(slide10, x + Inches(0.1), Inches(2.3), Inches(2.0), Inches(0.5),
                title, font_size=20, font_color=KDDI_WHITE, bold=True,
                alignment=PP_ALIGN.CENTER)
    add_textbox(slide10, x + Inches(0.1), Inches(2.85), Inches(2.0), Inches(1.2),
                desc, font_size=13, font_color=KDDI_WHITE,
                alignment=PP_ALIGN.CENTER)

# 総括
add_shape(slide10, Inches(0.6), Inches(4.5), Inches(12.1), Inches(0.05), KDDI_ORANGE)

summary_items = [
    "通信事業の安定収益基盤の上に、AI・DX・金融・エネルギーの非通信事業を積み上げる多角化戦略",
    "au Starlink Directは国内初の衛星直接通信として先行者利益を確保",
    "ローソン提携は短期的リスクはあるが、リアルテック融合の中長期的ポテンシャルは大きい",
    "営業利益1兆円超の安定基盤 + 持続的増配（配当性向40%超）で株主還元も充実",
    "Beyond 5G/6G時代の「社会を支えるプラットフォーマー」への変革が最大のテーマ",
]

add_bullet_points(slide10, Inches(0.8), Inches(4.7), Inches(11.5), Inches(2.5),
                  summary_items, font_size=14, bullet_char="◆", line_spacing=1.5)

# KDDI VISION
add_rounded_rect(slide10, Inches(2.5), Inches(6.7), Inches(8.3), Inches(0.6), KDDI_ORANGE)
add_textbox(slide10, Inches(2.7), Inches(6.73), Inches(7.9), Inches(0.5),
            "KDDI VISION 2030: 「つなぐチカラ」を進化させ、誰もが思いを実現できる社会をつくる。",
            font_size=15, font_color=KDDI_WHITE, bold=True, alignment=PP_ALIGN.CENTER)

add_page_number(slide10, 10, 10)


# =====================================================
# 保存
# =====================================================
output_path = "/home/user/test/KDDI_将来性_プレゼンテーション.pptx"
prs.save(output_path)
print(f"Presentation saved to: {output_path}")
